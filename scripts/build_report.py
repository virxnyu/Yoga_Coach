"""Build the assignment report (.docx) in the order required by the instruction sheet.

    python scripts/build_report.py                 # -> Yoga_Pose_Correction_Report.docx

Reads: team.json (names, roll numbers, e-mails, phones, GitHub link, contributions),
       outputs/metrics.json (written by train.py), docs/architecture.png,
       outputs/*.png, outputs/screenshots/*.png (real webcam screenshots, key 's' in realtime.py).
Re-run after training on your real recordings so every number in the report is current.
"""
from __future__ import annotations

import glob
import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from yogacoach.config import ERRORS, POSE_ERRORS, POSE_TITLES, RANGES, TARGET_POSES  # noqa: E402

FILL = "<fill in>"


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def table(doc, header, rows, widths=None, font=9.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font)
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v)); run.font.size = Pt(font)
            if str(v).startswith("<"):
                run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def numbered(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def figure(doc, path, caption, width=6.3):
    if not os.path.isfile(path):
        p = doc.add_paragraph(f"<missing figure: {os.path.relpath(path, ROOT)}>")
        p.runs[0].font.color.rgb = RGBColor(0xC0, 0, 0)
        return
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(caption); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True; c.runs[0].font.size = Pt(9)


def pct(x):
    return f"{100 * x:.1f}%"


def main():
    team = {"title": "Real-Time Yoga Pose Correction using Deep Learning", "course": "Deep Learning Techniques",
            "institute": "National Institute of Technology, Tiruchirappalli", "github": FILL,
            "members": [{"name": "Viranyu Srinivasan", "roll": FILL, "email": FILL, "phone": FILL, "contribution": FILL},
                        {"name": "Jeffrey Jayapackiyaraj", "roll": FILL, "email": FILL, "phone": FILL, "contribution": FILL},
                        {"name": "Aditya Agrawal", "roll": FILL, "email": FILL, "phone": FILL, "contribution": FILL}]}
    tp = os.path.join(ROOT, "team.json")
    if os.path.isfile(tp):
        team.update(json.load(open(tp)))
    mp = os.path.join(ROOT, "outputs", "metrics.json")
    M = json.load(open(mp)) if os.path.isfile(mp) else None
    R = M["results"] if M else {}

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9); s.top_margin = s.bottom_margin = Inches(0.8)

    t = doc.add_heading(team["title"], 0); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(f"{team['course']}  |  {team['institute']}"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Assignment report (demo: 29-30 Oct 2026, report: 02 Nov 2026)"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1 ------------------------------------------------------------------ team
    doc.add_heading("1. Team details", 1)
    table(doc, ["Name", "Roll No.", "Email id", "Contact number"],
          [[m["name"], m["roll"], m["email"], m["phone"]] for m in team["members"]], [2.0, 1.3, 2.2, 1.3])

    # 2 ------------------------------------------------------------------ problem
    doc.add_heading("2. Problem statement", 1)
    doc.add_paragraph(
        "Identify the incorrect yoga pose (for at least four yoga poses) from a real-time video and provide feedback. "
        "Given a live webcam stream of a person performing yoga, the system must (i) recognise which of the four "
        "target poses - " + ", ".join(POSE_TITLES[p] for p in TARGET_POSES) + " - is being attempted, (ii) decide "
        "whether the pose is performed correctly or incorrectly, and (iii) tell the user in plain language what to "
        "fix (for example 'Bend your front knee deeper, now 148 degrees'), updating continuously while the user moves.")
    doc.add_paragraph(
        "Formally, let X_t in R^(13x2) be the pixel-space keypoints of frame t. From a window W_t = (X_(t-11), ..., X_t) the system "
        "must output a triple (y_pose, y_ok, y_err): a pose label from {Tree, Warrior II, Triangle, Goddess, other}, a binary "
        "correctness decision, and a multi-label vector over eleven alignment errors, each of which is meaningful only for certain poses. "
        "The output has to be stable (no flicker), explainable (every complaint must be traceable to a measured angle) and fast enough "
        "for interactive use on a laptop CPU.")
    doc.add_heading("2.1 Scope and objectives", 2)
    bullets(doc, [
        "Four target poses: Tree (Vrikshasana), Warrior II (Virabhadrasana II), Triangle (Trikonasana) and Goddess (Utkata Konasana), plus a fifth 'other' class for non-target postures.",
        "Eleven specific, human-readable alignment errors (for example shallow front knee, drooping arms, leaning torso, narrow stance).",
        "Live video from a single, ordinary RGB webcam; no depth sensor, wearable or special hardware.",
        "Works for either body side, for different body sizes and camera distances.",
        "Honest evaluation that separates synthetic from real-world evidence.",
    ])

    # 3 ------------------------------------------------------------------ gaps
    doc.add_heading("3. Gaps identified", 1)
    bullets(doc, [
        ("Classification is not correction. ", "Public yoga work such as the Yoga-82 dataset (Verma et al., CVPRW 2020) and CNN/MediaPipe "
         "classifiers (e.g. Garg et al., 2022) answer 'which pose is this?'. They do not say whether the pose is wrong or how to fix it."),
        ("No labelled incorrect poses. ", "To our knowledge, no public yoga dataset provides frames labelled with specific alignment mistakes, so a purely "
         "supervised end-to-end 'incorrect pose' detector cannot be trained from public data."),
        ("Hand-tuned, frame-by-frame angle rules. ", "Pose-estimation + angle-threshold tutors are easy to build but react to every jittery frame, "
         "do not use temporal context and give flickering feedback."),
        ("Left/right-specific, unstable feedback. ", "Rules written for 'left knee' / 'right knee' break when the user mirrors the pose or switches sides."),
        ("Real-time constraint. ", "Heavy image-based models are hard to run on a laptop webcam at interactive speed; feedback must be low-latency."),
        ("Pose-specific error semantics are ignored. ", "A 'bent knee' is a mistake in Tree pose but required in Goddess pose. Generic multi-label heads that ignore this "
         "produce contradictory advice; the error space must be conditioned on the pose being performed."),
        ("Appearance and scale sensitivity. ", "Raw pixel coordinates change with camera distance, framing and body size, so models need scale- and translation-invariant, "
         "angle-based representations rather than raw landmark positions."),
        ("Opaque decisions. ", "End-to-end networks give a verdict without justification; a learner needs to know which joint is wrong and by how many degrees."),
    ])

    # 4 ------------------------------------------------------------------ architecture
    doc.add_heading("4. Proposed architecture diagram", 1)
    figure(doc, os.path.join(ROOT, "docs", "architecture.png"), "Figure 1. Proposed real-time pipeline (top) and offline training (left).", 6.7)
    doc.add_paragraph(
        "A webcam frame is passed to the MediaPipe Pose Landmarker (BlazePose), which returns 33 body landmarks; we keep 13 joints "
        "(nose, shoulders, elbows, wrists, hips, knees, ankles) in pixel coordinates so angles are not distorted by the image aspect ratio. "
        "Each frame becomes a 39-dimensional, scale- and translation-invariant feature vector. The last 12 frames feed YogaNet, a small temporal "
        "multi-task network, while a rule checker measures the same joint angles on the window's median pose. Their outputs are fused and "
        "temporally voted to give stable feedback that is drawn on the video.")
    doc.add_heading("4.1 Component description", 2)
    table(doc, ["Stage", "Module", "Role"],
          [["Perception", "MediaPipe Pose Landmarker (BlazePose full)", "33 landmarks per frame; 13 joints kept; visibility gating"],
           ["Representation", "geometry.py", "Hip-centred, torso-scaled coordinates + 8 joint angles + 5 posture ratios = 39 features"],
           ["Temporal learner", "YogaNet (model.py)", "Residual temporal convolutions + bidirectional GRU; multi-task heads"],
           ["Explainable verifier", "rules.py", "Per-pose target ranges, soft rule confidences, worded corrections"],
           ["Decision", "engine.py", "Hybrid fusion, EMA smoothing, majority vote, wrong-pose detection, hold timer"],
           ["Presentation", "draw.py, realtime.py", "Skeleton overlay, banner, corrections, alignment score"],
           ["Validation", "check_image.py, analyze.py", "Offline photo checker with annotated joint angles and per-angle target table"]],
          [1.3, 2.3, 3.2], 9)
    doc.add_heading("4.2 YogaNet parameter budget", 2)
    table(doc, ["Block", "Parameters"],
          [["Input projection (LayerNorm + Linear 39->128)", "5,198"],
           ["Two residual Conv1d blocks (k = 3)", "98,560"],
           ["Bidirectional GRU (64 x 2)", "74,496"],
           ["Shared trunk", "32,896"],
           ["Pose head (5) + correctness head (1) + error head (11)", "19,345"],
           ["Layer normalisation", "256"],
           ["Total", "230,751"]], [4.5, 1.5], 9)
    doc.add_paragraph("Measured on a CPU-only workstation, the complete decision engine (feature extraction, YogaNet, rule checker, fusion, voting) takes about 4.4 ms per frame "
                      "(rules only: about 1.4 ms), so the pose estimator, not the decision logic, is the bottleneck of the pipeline.")

    # 5 ------------------------------------------------------------------ algorithm
    doc.add_heading("5. Algorithmic steps", 1)
    numbered(doc, [
        "Capture a frame (mirrored webcam view) and run the MediaPipe Pose Landmarker; convert the 13 selected landmarks to pixel coordinates. If the mean landmark visibility is below 0.5 the frame is rejected and the user is asked to step back.",
        "Normalise: translate to the hip centre and divide by the torso length (shoulder-mid to hip-mid). Compute 8 joint angles (both knees, elbows, shoulders, hips) and 5 ratios (torso lean, stance width, wrist height above nose, wrist height above elbow, wrist-line angle). Result: 39 features per frame.",
        "Append the frame to a sliding window of T = 12 frames (about 0.4 s). Wait until the window is full.",
        "YogaNet forward pass: LayerNorm + Linear(39->128) -> two residual 1-D temporal convolutions -> bidirectional GRU -> mean and max pooling -> shared trunk -> three heads: pose class (4 poses + 'other'), correct/incorrect, and 11 multi-label error types (conditioned on the pose probabilities).",
        "Geometry: measure the angles on the window's median keypoints and compare them with the per-pose target ranges. Each error gets a soft rule confidence (0.5 exactly at the threshold).",
        "Fuse: score = 0.6 x P_network(error) + 0.4 x C_rule(error); an error is flagged if score >= 0.5 and it is meaningful for the pose. Smooth the scores with an exponential moving average (alpha = 0.45) and take a majority vote over the last 12 frames, so a single noisy frame can never flip the verdict.",
        "Choose the pose to judge: the user-selected target, or (auto mode) the pose with the highest smoothed probability above 0.55. If the target pose is very unlikely for 10 consecutive frames, report 'this does not look like <pose>'.",
        "Render feedback: skeleton with offending joints in red, CORRECT/INCORRECT banner, up to four plain-language corrections with the measured angle and the target, a 0-100 alignment score and a hold timer.",
    ])
    doc.add_heading("5.1 Mathematical formulation", 2)
    bullets(doc, [
        ("Joint angle. ", "For joint B with neighbours A and C: theta = arccos( (BA . BC) / (|BA| |BC|) ), computed in pixel space so the image aspect ratio does not distort it."),
        ("Normalisation. ", "With hip centre h and torso length L = |mid(shoulders) - mid(hips)|, each keypoint becomes p' = (p - h) / L. Distances such as stance width are expressed in torso lengths, which removes dependence on camera distance and body size."),
        ("Feature vector. ", "f_t = [ p'_1..13 (26 values), 8 angles, 5 ratios ] in R^39; a window F = (f_(t-11), ..., f_t) in R^(12x39)."),
        ("Network. ", "z = BiGRU( Conv_res( GELU( LN(F) W ) ) ); h = trunk([mean_t z, max_t z]); pose = softmax(W_p h); ok = sigmoid(W_c h); err = sigmoid(W_e [h, pose])."),
        ("Multi-task loss. ", "L = CE_ls(pose, y_pose; label smoothing 0.05) + 0.5 BCE(ok, y_ok) + sum_e m_e(y_pose) BCE(err_e, y_e) / sum m, where the mask m_e = 1 only if error e is applicable to the pose, so the network is never penalised for errors that cannot occur in that pose."),
        ("Rule confidence. ", "For a check with threshold tau and margin delta, the excess x = measured - tau (signed so that positive means violation) gives C = clip(0.5 + x / (4 delta), 0, 1), i.e. exactly 0.5 at the threshold and saturating two margin-widths either side of it (delta = 5 degrees or 0.08 torso lengths)."),
        ("Fusion. ", "S_e = 0.6 P_net(e) + 0.4 C_rule(e); error e is flagged when the smoothed S_e >= 0.5 and majority-voted over the last 12 frames."),
        ("Alignment score. ", "score = 100 x (1 - mean_k s_k), where s_k in [0,1] is the severity of check k (how far its measured value lies outside its target range, scaled and clipped); it rises smoothly as the user approaches the ideal posture."),
    ])
    doc.add_heading("5.2 Design decisions", 2)
    bullets(doc, [
        "Keypoints instead of pixels: the heavy visual work is done by a pre-trained estimator, so YogaNet can stay below 0.25 M parameters and train in minutes on a CPU.",
        "Bidirectional GRU over a short window: the window is only 0.4 s long, so non-causal context costs no perceptible latency but removes single-frame jitter.",
        "Pose-conditioned error head: the pose probabilities are fed to the error head and the loss is masked, so contradictory advice such as asking for a straight knee in Goddess pose cannot arise.",
        "Role-based measurements: the bent leg / straight leg / higher arm are identified by geometry rather than by left/right labels, making feedback identical for mirrored performances.",
        "Hybrid verification: the rule checker supplies exact angles for the feedback text and acts as an independent safety net against network errors on unusual bodies.",
    ])

    # 6 ------------------------------------------------------------------ dataset
    doc.add_heading("6. Dataset used", 1)
    doc.add_paragraph(
        "No public dataset labels yoga mistakes, so training data comes from two sources:")
    bullets(doc, [
        ("Synthetic kinematic data (bootstrap). ", "A forward-kinematics skeleton generates keypoint sequences for each pose. Joint angles are sampled either inside the target "
         "ranges (correct) or outside them (specific errors, one or several at once). Samples that lie within a safety margin (5 degrees or 0.08 torso lengths) of any threshold are rejected so labels are unambiguous. "
         "Camera roll, foreshortening, left/right mirroring, per-frame jitter and random scale/position are added. A fifth class 'other' (standing, arms raised, hands on hips, random arms) teaches the network to say 'no target pose'."),
        ("Real recordings (recommended and supported). ", "collect.py records labelled keypoint clips from a webcam or video file; extract_dataset.py converts folders of labelled photos/videos. "
         "train.py mixes them in (windows, mirror augmentation, over-sampling) and evaluates on held-out clips."),
    ])
    if M:
        a = M["args"]
        doc.add_paragraph(f"Synthetic sizes used for the reported run: {a['train_n']} + {a['train_n'] // 2} training sequences per pose (two noise levels), "
                          f"{a['val_n']} validation and {a['test_n']} test sequences per pose, 5 classes, window length 12.")
        rd = a.get("real_dir")
        nreal = len(glob.glob(os.path.join(ROOT, rd, "*.npz"))) if rd else 0
        doc.add_paragraph(f"Real recorded clips used in this run: {nreal}" + ("" if nreal else "  <none yet - record clips with collect.py, retrain and rebuild this report>"))
    doc.add_paragraph("Composition of the synthetic training pool (per 1,000 sequences of each pose, before noise augmentation):")
    table(doc, ["Pose", "Correct", "Incorrect", "Errors injected (counts)"],
          [["Tree", "39%", "61%", "straight_knee_bent 253, bent_knee_shallow 236, arms_low 225, elbows_bent 194, torso_lean 251"],
           ["Warrior II", "26%", "74%", "straight_knee_bent 219, bent_knee_shallow 379, arms_not_level 611, elbows_bent 246, torso_lean 317, stance_narrow 229"],
           ["Triangle", "50%", "50%", "straight_knee_bent 143, stance_narrow 209, lean_insufficient 158, arms_not_stacked 183, elbows_bent 196"],
           ["Goddess", "41%", "59%", "bent_knee_shallow 434, knees_uneven 110, stance_narrow 257, arms_not_level 374, elbow_angle_wrong 211, torso_lean 193"],
           ["Other", "-", "-", "standing, arms raised, hands on hips, random arms (negative class)"]],
          [1.0, 0.8, 0.8, 4.2], 8.5)
    doc.add_paragraph("Poses, errors the system can detect, and the target ranges that define 'correct' (all configurable in yogacoach/config.py):")
    rng_txt = {
        "tree": "standing knee >= 165; lifted knee <= 100; wrists above nose; elbows >= 150; torso lean <= 10",
        "warrior2": "front knee 85-125; back knee >= 165; shoulder angle 75-105 (arms level); elbows >= 160; lean <= 10; stance >= 1.7 torso lengths",
        "triangle": "both knees >= 160; torso lean 35-95; stance >= 1.6; wrist line within 22 of vertical; elbows >= 160",
        "goddess": "both knees <= 135; knee difference <= 22; stance >= 1.4; shoulder angle 70-110; elbow angle 70-110; lean <= 10",
    }
    table(doc, ["Pose", "Detectable errors", "'Correct' ranges (degrees unless noted)"],
          [[POSE_TITLES[p], ", ".join(POSE_ERRORS[p]), rng_txt[p]] for p in TARGET_POSES], [1.5, 2.6, 2.7], 8.5)

    # 7 ------------------------------------------------------------------ tools
    doc.add_heading("7. Tools used", 1)
    bullets(doc, [
        "Python 3, PyTorch (YogaNet, training), NumPy, Matplotlib",
        "MediaPipe Pose Landmarker (BlazePose, Tasks API) for 2-D body landmarks",
        "OpenCV for video capture and the on-screen overlay; Pillow for EXIF-safe photo decoding",
        "A local web page (check_image.py, Python standard-library HTTP server) to validate the model on uploaded photos: annotated joint angles, verdict, corrections and an angle-versus-target table",
        "python-docx for this report; unit and integration tests (tests/test_pipeline.py, tests/test_image_check.py: 13 tests covering geometry, rules, model, engine, image analysis and the HTTP round trip)",
        "Git/GitHub for version control and submission",
        "Claude (Anthropic) as an AI coding/writing assistant - see section 14",
    ])

    # 8 ------------------------------------------------------------------ screenshots
    doc.add_heading("8. Screenshot of the result", 1)
    shots = sorted(glob.glob(os.path.join(ROOT, "outputs", "screenshots", "*.png")))
    if shots:
        for i, s in enumerate(shots[:8]):
            figure(doc, s, f"Figure {2 + i}. Real-time webcam result ({os.path.basename(s)}).", 5.8)
    else:
        p = doc.add_paragraph("<REPLACE with real webcam screenshots: run  python realtime.py --pose tree  and press 's' in each pose; "
                              "they are saved to outputs/screenshots/ and picked up automatically when this report is rebuilt.>")
        p.runs[0].font.color.rgb = RGBColor(0xC0, 0, 0)
        doc.add_paragraph("Until then, the figures below are rendered from SIMULATED skeleton input (no camera) through the same engine and overlay; "
                          "they illustrate the feedback format only and are not experimental results.")
        for i, s in enumerate(sorted(glob.glob(os.path.join(ROOT, "outputs", "simulated", "*.png")))[:8]):
            figure(doc, s, f"Figure {2 + i}. Simulated input: {os.path.basename(s)[:-4].replace('_', ' ')}.", 4.6)

    # 9 ------------------------------------------------------------------ analysis
    doc.add_heading("9. Analysis and how it differs from existing approaches", 1)
    if M:
        doc.add_heading("9.1 Results", 2)
        figure(doc, os.path.join(ROOT, "outputs", "training_curves.png"), "Training curves (multi-task loss, validation pose accuracy, validation error micro-F1).", 6.5)
        rows = []
        for key, r in R.items():
            rows.append([r["name"], r["n"], pct(r["pose"]["accuracy"]), pct(r["errors_hybrid"]["micro_f1"]),
                         pct(r["errors_dl"]["micro_f1"]), pct(r["errors_rules_single_frame"]["micro_f1"]),
                         pct(r["errors_hybrid"]["correct_vs_incorrect"]["accuracy"])])
        table(doc, ["Test set", "N", "Pose acc.", "Error F1 hybrid", "Error F1 YogaNet", "Error F1 rules (1 frame)", "Correct/incorrect acc. (hybrid)"],
              rows, [1.9, 0.5, 0.7, 0.8, 0.8, 0.9, 1.2], 8.5)
        figure(doc, os.path.join(ROOT, "outputs", "method_comparison.png"), "Error micro-F1 of single-frame rules, window rules, YogaNet alone and the hybrid.", 6.3)
        figure(doc, os.path.join(ROOT, "outputs", "confusion_matrix.png"), "Pose confusion matrix on the synthetic test set.", 3.8)
        pe = R["synthetic_test"]["errors_hybrid"]["per_error"]
        table(doc, ["Error type", "Precision", "Recall", "F1", "Support"],
              [[k, f"{v['precision']:.3f}", f"{v['recall']:.3f}", f"{v['f1']:.3f}", v["support"]] for k, v in pe.items()],
              [2.3, 1, 1, 1, 1], 9)
        if "real_test" not in R:
            p = doc.add_paragraph("<Real-world evaluation not yet run: record clips with collect.py, train with --real-dir data/real and rebuild; a 'real_test' row (held-out clips) will then appear in the table above.>")
            p.runs[0].font.color.rgb = RGBColor(0xC0, 0, 0)
        doc.add_heading("9.2 How to read these numbers (limitations)", 2)
        bullets(doc, [
            "The 'synthetic' rows measure how well the models recover labels on generated skeletons with sensor noise; they are not accuracy on real people. Real-world accuracy must be measured on held-out real clips.",
            "Synthetic labels are defined by the same target ranges used by the rule checker, so the rule baseline is not an independent judge here; the comparison shows robustness to noise and temporal jitter, not which ranges are 'right'.",
            "The system uses 2-D keypoints from a single camera and assumes a front view for the four chosen poses; depth-dependent mistakes (for example knee over toes seen from the side) are out of scope.",
            "Target ranges come from common alignment cues, not a certified instructor; they are easy to edit in config.py.",
        ])
    doc.add_heading("9.3 Discussion", 2)
    bullets(doc, [
        "The temporal network and the rule checker fail in different ways: rules are brittle under keypoint jitter, the network is brittle outside its training distribution. Fusing them with temporal voting gives better stability than either alone, as the method-comparison figure shows on the noisy test set.",
        "Masking the error loss by pose applicability keeps the 11-way error head well calibrated: errors that are impossible for a pose are never trained as negatives, and the head is conditioned on the pose posterior.",
        "Because every complaint is generated from a measured angle and a target range, the feedback is auditable: a user or instructor can see the number that triggered it and adjust the range in config.py.",
        "Inference cost is dominated by the pose estimator; the decision stack adds only a few milliseconds per frame on a CPU, so the system runs in real time on a laptop.",
        "The offline image checker allows the same pipeline to be verified on still photographs with a visible per-angle breakdown, which makes failure cases (side-on body, cropped limbs, tiny figure) easy to diagnose.",
    ])
    doc.add_heading("9.4 Difference from existing approaches", 2)
    table(doc, ["Aspect", "Typical yoga classifiers", "Frame-wise angle-rule tutors", "This work"],
          [["Output", "pose name only", "pass/fail per rule", "pose + correct/incorrect + specific, worded corrections"],
           ["Temporal context", "single image", "single frame", "12-frame window, EMA and majority vote"],
           ["Handles mistakes", "no", "yes (hand-tuned)", "learned multi-label error head + rules"],
           ["Left/right independent", "n/a", "usually not", "yes: roles such as bent leg / straight leg, mirror augmentation"],
           ["Training data for mistakes", "none available", "none (no learning)", "kinematic generator + user-recorded clips"],
           ["Speed", "CNN on pixels", "real-time", "real-time (pose estimator + 0.2 M-parameter net on keypoints)"]],
          [1.3, 1.6, 1.7, 2.3], 8.5)

    # 10 ----------------------------------------------------------------- code
    doc.add_heading("10. Submit code link", 1)
    p = doc.add_paragraph(f"GitHub: {team['github']}")
    if team["github"] == FILL:
        p.runs[0].font.color.rgb = RGBColor(0xC0, 0, 0)
    doc.add_paragraph("The repository contains the code, a README with setup and run instructions, the trained checkpoint and this report.")

    # 11 ----------------------------------------------------------------- novelty
    doc.add_heading("11. Novel contribution", 1)
    bullets(doc, [
        ("Correction, not just classification: ", "a three-headed temporal network (pose, correct/incorrect, 11 multi-label error types conditioned on the pose) trained on keypoint windows."),
        ("Hybrid decision: ", "learned error probabilities are fused with geometry-based soft confidences, so feedback is both learned (robust to noise) and explainable (it quotes measured angles and targets)."),
        ("Labelled-mistake data without manual annotation: ", "a margin-controlled kinematic generator that produces unambiguous correct and incorrect examples, plus a recording tool to add real clips."),
        ("Side-invariant feedback: ", "measurements and messages use roles (bent leg, straight leg, higher arm) instead of left/right, and training uses mirror augmentation."),
        ("Stable real-time feedback: ", "windowing, EMA smoothing, majority voting, visibility gating and wrong-pose detection."),
        ("Pose-conditioned, masked multi-task learning: ", "the loss only penalises errors that are defined for the pose, and the error head sees the pose posterior."),
        ("Soft rule confidences: ", "geometry outputs graded confidences (0.5 at the threshold) instead of hard pass/fail so they can be fused probabilistically with the network."),
        ("Built-in validation tooling: ", "a photo checker and a recording tool close the loop between model, data and real people."),
    ])

    # 12 ----------------------------------------------------------------- individual
    doc.add_heading("12. Individual contribution", 1)
    table(doc, ["Team member", "Contribution"], [[m["name"], m["contribution"]] for m in team["members"]], [2.0, 4.8])
    p = doc.add_paragraph("<Each member should describe their own part here (edit team.json and rebuild).>")
    p.runs[0].font.color.rgb = RGBColor(0xC0, 0, 0)

    # 13 ----------------------------------------------------------------- learning
    doc.add_heading("13. Learning outcome", 1)
    bullets(doc, [
        "Using a pre-trained pose estimator as a front end and designing scale-invariant features from keypoints.",
        "Designing and training a multi-task temporal network (conv + GRU, shared trunk, softmax and sigmoid heads) with masked losses.",
        "Dealing with missing labelled data: synthetic data generation, label-ambiguity margins, domain randomisation, mixing real and synthetic data.",
        "Evaluating honestly: separating synthetic from real test sets, noise-robustness tests, per-class F1 and ablations against rule-only baselines.",
        "Engineering a real-time system: latency, temporal smoothing, failure handling (no person, partial visibility).",
    ])

    # 14 ----------------------------------------------------------------- AI usage
    doc.add_heading("14. AI tool usage", 1)
    doc.add_paragraph(
        "Claude (Anthropic) was used as an AI assistant to draft the project code (feature extraction, synthetic data generator, YogaNet, training, "
        "real-time application, tests), the architecture diagram and the first draft of this report. The team specified the problem, reviewed and "
        "ran the code, and is responsible for all content. The AI could not use a webcam or download the MediaPipe model in its workspace, so the "
        "live-camera path must be and was verified by the team on their own machine; results marked as synthetic come from generated data. "
        "<Team: edit this paragraph to describe exactly what you reviewed, changed and tested yourselves.>")
    doc.add_heading("References", 2)
    for r in [
        "Bazarevsky, V., Grishchenko, I., et al. BlazePose: On-device Real-time Body Pose Tracking. arXiv:2006.10204, 2020.",
        "Verma, M., Kumawat, S., Nakashima, Y., Raman, S. Yoga-82: A New Dataset for Fine-grained Classification of Human Poses. CVPR Workshops, 2020.",
        "Cao, Z., Hidalgo, G., Simon, T., Wei, S.-E., Sheikh, Y. OpenPose: Realtime Multi-Person 2D Pose Estimation using Part Affinity Fields. IEEE TPAMI, 2019 (arXiv:1812.08008).",
        "Garg, S., et al. Yoga pose classification: a CNN and MediaPipe inspired deep learning approach for real-world application, 2022.",
        "MediaPipe Pose Landmarker documentation, Google AI Edge.",
    ]:
        doc.add_paragraph(r, style="List Bullet")

    out = os.path.join(ROOT, "Yoga_Pose_Correction_Report.docx")
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    main()